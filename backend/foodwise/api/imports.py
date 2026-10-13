"""Bring-your-own-data onboarding: upload, map columns, match dishes, dry-run report, commit."""

from __future__ import annotations

import csv
import io
import uuid
from datetime import UTC, datetime

from fastapi import APIRouter, Depends, File, UploadFile
from fastapi.responses import PlainTextResponse
from pydantic import BaseModel, Field
from sqlalchemy import func, select
from sqlalchemy.dialects.postgresql import insert

from foodwise.core import audit, jobs
from foodwise.core.deps import Ctx, need
from foodwise.core.errors import Conflict, NotFound, Unprocessable
from foodwise.core.storage import storage, tenant_key
from foodwise.db.models import CalendarDay, Canteen, Dish, ImportJob, ServiceLog, Tenant, WeatherDay
from foodwise.domain import dataquality as dq
from foodwise.domain.readiness import readiness
from foodwise.domain.settings import effective_settings

router = APIRouter(prefix="/imports", tags=["imports"])

CANONICAL = ["date", "canteen", "meal", "dish", "unit", "prepared_qty", "served_qty", "leftover_qty", "footfall",
             "sold_out", "weather", "event"]
REQUIRED = ["date", "meal", "dish", "prepared_qty", "served_qty"]
TEMPLATE = (
    "date,canteen,meal,dish,unit,prepared_qty,served_qty,leftover_qty,footfall,sold_out,weather,event\n"
    "2026-08-18,North Mess,Lunch,Vegetable pulao,plates,350,331,19,1020,false,Clear,\n"
    "2026-08-19,North Mess,Lunch,Vegetable pulao,plates,360,360,0,1180,true,Rain,Guest lecture\n"
)
SYNONYMS = {
    "date": ["date", "day", "service date", "dt"],
    "canteen": ["canteen", "mess", "outlet", "kitchen", "location"],
    "meal": ["meal", "service", "session", "shift"],
    "dish": ["dish", "item", "menu item", "food", "dish name"],
    "unit": ["unit", "uom", "units"],
    "prepared_qty": ["prepared", "prepared qty", "cooked", "made", "produced", "prepared quantity"],
    "served_qty": ["served", "served qty", "consumed", "sold", "issued", "served quantity"],
    "leftover_qty": ["leftover", "left over", "wasted", "waste", "remaining", "unserved"],
    "footfall": ["footfall", "diners", "headcount", "attendance", "covers"],
    "sold_out": ["sold out", "soldout", "ran out", "stockout"],
    "weather": ["weather", "condition"],
    "event": ["event", "notes", "occasion"],
}
MAX_BYTES = 10 * 1024 * 1024
MAX_ROWS = 100_000


class ImportOut(BaseModel):
    id: uuid.UUID
    filename: str
    status: str
    headers: list[str]
    preview: list[list[str]]
    mapping: dict[str, str | None]
    row_count: int
    report: dict
    committed_rows: int
    created_at: datetime
    committed_at: datetime | None


def _out(j: ImportJob) -> ImportOut:
    return ImportOut(id=j.id, filename=j.filename, status=j.status, headers=j.headers, preview=j.preview,
                     mapping=j.mapping, row_count=j.row_count, report=j.report, committed_rows=j.committed_rows,
                     created_at=j.created_at, committed_at=j.committed_at)


def _read_rows(data: bytes, filename: str) -> list[list[str]]:
    if filename.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        rows = []
        for row in ws.iter_rows(values_only=True):
            rows.append(["" if v is None else (v.date().isoformat() if isinstance(v, datetime) else str(v)) for v in row])
            if len(rows) > MAX_ROWS + 1:
                break
        return rows
    text = data.decode("utf-8-sig", errors="replace")
    dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t") if text.strip() else csv.excel
    return [r for r in csv.reader(io.StringIO(text), dialect) if any(c.strip() for c in r)]


def _suggest(headers: list[str]) -> dict[str, str | None]:
    mapping: dict[str, str | None] = {}
    used: set[str] = set()
    for field in CANONICAL:
        best, score = None, 0.0
        for h in headers:
            if h in used:
                continue
            s = max(dq.similarity(h, syn) for syn in SYNONYMS[field])
            if s > score:
                best, score = h, s
        mapping[field] = best if score >= 0.75 else None
        if mapping[field]:
            used.add(mapping[field])
    return mapping


@router.get("/template.csv", response_class=PlainTextResponse)
def template() -> PlainTextResponse:
    return PlainTextResponse(TEMPLATE, media_type="text/csv",
                             headers={"Content-Disposition": "attachment; filename=foodwise_import_template.csv"})


@router.get("", response_model=list[ImportOut])
def list_imports(ctx: Ctx = Depends(need("imports:write"))) -> list[ImportOut]:
    return [_out(j) for j in ctx.db.execute(select(ImportJob).order_by(ImportJob.created_at.desc()).limit(50)).scalars()]


@router.post("", response_model=ImportOut, status_code=201)
async def upload(file: UploadFile = File(...), ctx: Ctx = Depends(need("imports:write"))) -> ImportOut:
    name = (file.filename or "upload.csv")[:200]
    if not name.lower().endswith((".csv", ".txt", ".xlsx", ".xlsm")):
        raise Unprocessable("Upload a CSV or Excel (.xlsx) file.")
    data = await file.read(MAX_BYTES + 1)
    if len(data) > MAX_BYTES:
        raise Unprocessable("That file is larger than 10 MB. Split it by year and upload the parts.")
    if name.lower().endswith((".xlsx", ".xlsm")) and not data.startswith(b"PK"):
        raise Unprocessable("That file does not look like an Excel workbook.")
    try:
        rows = _read_rows(data, name)
    except Exception as exc:  # noqa: BLE001
        raise Unprocessable("We could not read that file. Save it as CSV (UTF-8) and try again.") from exc
    if len(rows) < 2:
        raise Unprocessable("The file needs a header row and at least one data row.")
    if len(rows) > MAX_ROWS + 1:
        raise Unprocessable(f"The file has more than {MAX_ROWS:,} rows. Split it and upload the parts.")
    headers = [h.strip() for h in rows[0]]
    job_id = uuid.uuid4()
    key = tenant_key(ctx.tenant_id, "imports", f"{job_id}{'.xlsx' if name.lower().endswith(('.xlsx', '.xlsm')) else '.csv'}")
    storage().put(key, data)
    j = ImportJob(id=job_id, tenant_id=ctx.tenant_id, filename=name, storage_key=key, status="uploaded",
                  headers=headers, preview=rows[1:9], mapping=_suggest(headers), row_count=len(rows) - 1,
                  created_by=ctx.principal.user_id)
    ctx.db.add(j)
    audit.record(ctx.db, ctx.principal, "import.uploaded", target_type="import", target_id=job_id,
                 meta={"filename": name, "rows": len(rows) - 1})
    ctx.db.commit()
    return _out(j)


class ValidateIn(BaseModel):
    mapping: dict[str, str | None]
    default_canteen_id: uuid.UUID | None = None
    dish_matches: dict[str, str] = Field(default_factory=dict, description="raw dish name -> dish id, or 'new'")


def _parse(ctx: Ctx, j: ImportJob, body: ValidateIn) -> tuple[list[dict], dict, dict, list]:
    missing = [f for f in REQUIRED if not body.mapping.get(f)]
    if missing:
        raise Unprocessable(f"Map these columns first: {', '.join(missing)}.")
    rows = _read_rows(storage().get(j.storage_key), j.storage_key)
    header = [h.strip() for h in rows[0]]
    idx = {f: header.index(col) for f, col in body.mapping.items() if col and col in header}
    canteens = ctx.db.execute(select(Canteen)).scalars().all()
    if ctx.principal.canteen_ids:
        canteens = [c for c in canteens if c.id in ctx.principal.canteen_ids]
    if not canteens:
        raise Unprocessable("Create a canteen before importing its history.")
    default = next((c for c in canteens if c.id == body.default_canteen_id), canteens[0] if len(canteens) == 1 else None)
    dishes = ctx.db.execute(select(Dish)).scalars().all()
    dish_list = [(str(d.id), d.name, list(d.aliases or [])) for d in dishes]
    dish_units = {str(d.id): d.unit for d in dishes}

    def cell(r: list[str], f: str) -> str:
        i = idx.get(f)
        return r[i].strip() if i is not None and i < len(r) else ""

    matches: dict[str, dict] = {}
    canteen_matches: dict[str, str | None] = {}
    parsed = []
    for n, r in enumerate(rows[1:], start=2):
        raw_c = cell(r, "canteen")
        if raw_c not in canteen_matches:
            best = max(canteens, key=lambda c: dq.similarity(raw_c, c.name)) if raw_c else default
            ok = best is not None and (not raw_c or dq.similarity(raw_c, best.name) >= 0.75)
            canteen_matches[raw_c] = str(best.id) if ok else (str(default.id) if default else None)
        canteen_id = canteen_matches[raw_c]
        canteen = next((c for c in canteens if str(c.id) == canteen_id), None)
        raw_meal = cell(r, "meal")
        meal = None
        if canteen is not None:
            for m in canteen.meals:
                if raw_meal.lower() in (m["key"], m["label"].lower()) or dq.similarity(raw_meal, m["label"]) >= 0.8:
                    meal = m["key"]
        raw_dish = cell(r, "dish")
        if raw_dish not in matches:
            chosen = body.dish_matches.get(raw_dish)
            if chosen and chosen != "new":
                d = next((x for x in dishes if str(x.id) == chosen), None)
                matches[raw_dish] = {"raw": raw_dish, "dish_id": chosen, "dish_name": d.name if d else None,
                                     "score": 1.0, "status": "confirmed"}
            elif chosen == "new":
                matches[raw_dish] = {"raw": raw_dish, "dish_id": None, "dish_name": None, "score": 0.0, "status": "new"}
            else:
                m = dq.match_dish(raw_dish, dish_list)
                matches[raw_dish] = {"raw": raw_dish, "dish_id": m.dish_id, "dish_name": m.dish_name, "score": m.score,
                                     "status": m.status}
        match = matches[raw_dish]
        dish_id = match["dish_id"] if match["status"] in ("exact", "confirmed", "suggested") else None
        if match["status"] == "new":
            dish_id = f"new:{raw_dish}"
        leftover = dq.parse_float(cell(r, "leftover_qty"))
        parsed.append({
            "_row": n, "date": dq.parse_date(cell(r, "date")), "canteen_id": canteen_id, "meal": meal, "meal_raw": raw_meal,
            "dish_id": dish_id, "dish_raw": raw_dish, "unit": cell(r, "unit").lower() or None,
            "prepared_qty": dq.parse_float(cell(r, "prepared_qty")), "served_qty": dq.parse_float(cell(r, "served_qty")),
            "leftover_qty": leftover if leftover is not None else None,
            "footfall": dq.parse_float(cell(r, "footfall")), "sold_out": dq.parse_bool(cell(r, "sold_out")),
            "weather": cell(r, "weather"), "event": cell(r, "event"),
        })
    for p in parsed:
        if p["leftover_qty"] is None and p["prepared_qty"] is not None and p["served_qty"] is not None:
            p["leftover_qty"] = max(0.0, p["prepared_qty"] - p["served_qty"])
    known_meals = {m["key"] for c in canteens for m in c.meals}
    return parsed, matches, canteen_matches, [known_meals, dish_units]


@router.post("/{import_id}/validate", response_model=ImportOut)
def validate(import_id: uuid.UUID, body: ValidateIn, ctx: Ctx = Depends(need("imports:write"))) -> ImportOut:
    j = ctx.db.get(ImportJob, import_id)
    if j is None:
        raise NotFound("Import not found.")
    if j.status == "committed":
        raise Conflict("This import is already committed.")
    parsed, matches, canteen_matches, (known_meals, dish_units) = _parse(ctx, j, body)
    report = dq.check_rows(parsed, known_meals=known_meals, dish_units=dish_units)
    tenant = ctx.db.get(Tenant, ctx.tenant_id)
    existing_days = {
        str(cid): n for cid, n in ctx.db.execute(
            select(ServiceLog.canteen_id, func.count(func.distinct(ServiceLog.service_date))).group_by(ServiceLog.canteen_id)
        ).all()
    }
    new_days: dict[str, set] = {}
    for p in parsed:
        if p["date"] and p["canteen_id"]:
            new_days.setdefault(p["canteen_id"], set()).add(p["date"])
    tiers = effective_settings(tenant.settings)["tiers"]
    readiness_after = {}
    for cid, days in new_days.items():
        r = readiness(existing_days.get(cid, 0) + len(days), tiers)
        readiness_after[cid] = {"tier": r.tier, "name": r.name, "days": r.days, "message": r.message}
    j.mapping = body.mapping
    j.dish_matches = {k: v for k, v in body.dish_matches.items()}
    j.canteen_matches = {k: v for k, v in canteen_matches.items() if v}
    j.status = "validated"
    j.report = {
        "score": report.score,
        "rows": report.rows,
        "valid_rows": report.valid_rows,
        "counts": report.counts(),
        "issues": [i.__dict__ for i in report.issues[:300]],
        "issues_truncated": len(report.issues) > 300,
        "date_range": report.date_range,
        "days": report.days,
        "gaps": report.gaps[:20],
        "dish_matches": sorted(matches.values(), key=lambda m: (m["status"] != "suggested", m["status"] != "new", m["raw"])),
        "canteen_matches": canteen_matches,
        "readiness_after": readiness_after,
        "default_canteen_id": str(body.default_canteen_id) if body.default_canteen_id else None,
    }
    ctx.db.commit()
    return _out(j)


class CommitIn(BaseModel):
    create_new_dishes: bool = True


@router.post("/{import_id}/commit", response_model=ImportOut)
def commit(import_id: uuid.UUID, body: CommitIn, ctx: Ctx = Depends(need("imports:write"))) -> ImportOut:
    j = ctx.db.get(ImportJob, import_id)
    if j is None:
        raise NotFound("Import not found.")
    if j.status != "validated":
        raise Conflict("Validate the import before committing it.")
    vin = ValidateIn(mapping=j.mapping, dish_matches=j.dish_matches,
                     default_canteen_id=j.report.get("default_canteen_id"))
    parsed, matches, _, (known_meals, dish_units) = _parse(ctx, j, vin)
    created: dict[str, uuid.UUID] = {}
    if body.create_new_dishes:
        for raw, m in matches.items():
            if m["status"] == "new" and raw:
                unit = next((p["unit"] for p in parsed if p["dish_raw"] == raw and p["unit"]), "plates")
                d = Dish(tenant_id=ctx.tenant_id, name=raw[:120], course="main",
                         unit=unit if unit in ("plates", "kg", "litres", "pieces") else "plates")
                ctx.db.add(d)
                ctx.db.flush()
                jobs.enqueue(ctx.db, "find_dish_image", {"dish_id": str(d.id)}, tenant_id=ctx.tenant_id,
                             dedupe_key=f"dish-image:{d.id}")
                created[f"new:{raw}"] = d.id
                dish_units[str(d.id)] = d.unit
    for p in parsed:
        if p["dish_id"] in created:
            p["dish_id"] = str(created[p["dish_id"]])
    report = dq.check_rows(parsed, known_meals=known_meals, dish_units=dish_units)
    bad = {i.row for i in report.issues if i.level == "error"}
    values = []
    weather, events = {}, set()
    for p in parsed:
        if p["_row"] in bad or p["date"] is None or not p["dish_id"] or str(p["dish_id"]).startswith("new:"):
            continue
        values.append({
            "id": uuid.uuid4(), "tenant_id": ctx.tenant_id, "canteen_id": uuid.UUID(p["canteen_id"]),
            "service_date": p["date"], "meal": p["meal"], "dish_id": uuid.UUID(str(p["dish_id"])),
            "unit": p["unit"] or dish_units.get(str(p["dish_id"]), "plates"), "prepared_qty": p["prepared_qty"],
            "served_qty": p["served_qty"], "unserved_qty": p["leftover_qty"] or 0.0, "plate_waste_qty": 0.0,
            "spoilage_qty": 0.0, "footfall": int(p["footfall"]) if p["footfall"] is not None else None,
            "sold_out": p["sold_out"], "source": "import", "import_job_id": j.id, "quality_flags": [],
        })
        if p["weather"]:
            weather[(p["canteen_id"], p["date"])] = p["weather"].lower()
        if p["event"]:
            events.add((p["date"], p["event"][:200]))
    for i in range(0, len(values), 2000):
        stmt = insert(ServiceLog).values(values[i : i + 2000])
        stmt = stmt.on_conflict_do_update(
            index_elements=["tenant_id", "canteen_id", "service_date", "meal", "dish_id"],
            set_={c: stmt.excluded[c] for c in ("prepared_qty", "served_qty", "unserved_qty", "footfall", "sold_out",
                                                 "unit", "import_job_id")},
        )
        ctx.db.execute(stmt)
    for (cid, day), cond in weather.items():
        rain = 0.8 if "rain" in cond or "storm" in cond else 0.35 if "cloud" in cond else 0.1
        stmt = insert(WeatherDay).values(id=uuid.uuid4(), tenant_id=ctx.tenant_id, canteen_id=uuid.UUID(cid), day=day,
                                         temp_c=27.0, rain_prob=rain, condition=cond[:20], source="import")
        ctx.db.execute(stmt.on_conflict_do_nothing())
    for day, label in events:
        ctx.db.add(CalendarDay(tenant_id=ctx.tenant_id, day=day, kind="event", label=label))
    j.status = "committed"
    j.committed_rows = len(values)
    j.committed_at = datetime.now(UTC)
    j.report = {**j.report, "committed_rows": len(values), "skipped_rows": len(parsed) - len(values),
                "new_dishes": len(created)}
    for cid in {v["canteen_id"] for v in values}:
        jobs.enqueue(ctx.db, "train_canteen", {"canteen_id": str(cid), "trigger": "import"}, tenant_id=ctx.tenant_id,
                     dedupe_key=f"train:{cid}")
    audit.record(ctx.db, ctx.principal, "import.committed", target_type="import", target_id=j.id,
                 meta={"rows": len(values), "new_dishes": len(created)})
    ctx.db.commit()
    return _out(j)


@router.get("/{import_id}", response_model=ImportOut)
def get_import(import_id: uuid.UUID, ctx: Ctx = Depends(need("imports:write"))) -> ImportOut:
    j = ctx.db.get(ImportJob, import_id)
    if j is None:
        raise NotFound("Import not found.")
    return _out(j)
